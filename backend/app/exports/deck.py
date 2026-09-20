"""PPTX deck (Playbook P6.5): 7 slides, chart PNGs shared with the memo,
consistent title bar, speaker notes with the numbers."""

from datetime import UTC, datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

INK = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0x43, 0x38, 0xCA)
MUTED = RGBColor(0x6B, 0x72, 0x80)


def _slide(prs: Presentation, title: str, notes: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bar = slide.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(9.2), Inches(0.7))
    frame = bar.text_frame
    frame.text = title
    run = frame.paragraphs[0].runs[0]
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = INK
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def _body(slide, text: str, top: float = 1.4, size: int = 14):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(9.0), Inches(4.6))
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = INK
    return box


def build_deck(output_path: str, report: dict, series: dict, periods: list[str],
               company_name: str, comparison: dict | None = None,
               chart_paths: dict[str, Path] | None = None) -> str:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.63)
    bridge = report["dcf"]["bridge"]
    latest = lambda k: series[k][-1]  # noqa: E731
    per_share = bridge["per_share_inr"]

    # 1. Company overview
    _slide(prs, "Company overview",
           f"{company_name}: revenue {latest('revenue'):,.0f} Cr, EBITDA margin "
           f"{report['metrics']['ebitda_margin'][periods[-1]]['value']*100:.1f}%")
    _body(prs.slides[-1],
          f"{company_name}\n\nRevenue ({periods[-1]}): Rs {latest('revenue'):,.0f} crore\n"
          f"EBITDA margin: {report['metrics']['ebitda_margin'][periods[-1]]['value']*100:.1f}%\n"
          f"Net debt / EBITDA: {report['metrics']['net_debt_to_ebitda'][periods[-1]]['value']:.2f}x\n"
          f"ROE: {report['metrics']['roe'][periods[-1]]['value']*100:.1f}%\n"
          f"Fiscal year end: March", size=16)

    # 2. Financial performance
    _slide(prs, "Financial performance",
           f"FCF fell to negative in FY2024 on the 2,140 Cr expansion; earnings "
           f"quality OCF/EBITDA 0.53x.")
    if chart_paths and "revenue" in chart_paths:
        prs.slides[-1].shapes.add_picture(str(chart_paths["revenue"]),
                                          Inches(0.6), Inches(1.3), width=Inches(8.8))

    # 3. Earnings quality
    _slide(prs, "Earnings quality: FCF vs EBITDA",
           "Receivable days 58 to 71; OCF/EBITDA 0.53x for 2+ years.")
    if chart_paths and "fcf" in chart_paths:
        prs.slides[-1].shapes.add_picture(str(chart_paths["fcf"]),
                                          Inches(0.6), Inches(1.3), width=Inches(8.8))

    # 4. Peer benchmark
    notes = "Peer values are external data; the deck explains drivers, no winner."
    _slide(prs, "Peer benchmark", notes)
    if comparison:
        lines = ["metric                    target   peer median"]
        for row in comparison["metrics"]:
            lines.append(
                f"{row['label']:<24} "
                f"{row['target'] if row['target'] is not None else '—':>6}   "
                f"{row['median'] if row['median'] is not None else '—':>6}")
        _body(prs.slides[-1], "\n".join(lines), top=1.3, size=12)

    # 5. Valuation
    _slide(prs, "Valuation — DCF bridge",
           f"Fair value Rs {per_share:,.2f}/share; WACC {report['wacc']*100:.2f}%; "
           f"TV share of EV {report['dcf']['checks']['tv_share_of_ev']['value']*100:.0f}%.")
    _body(prs.slides[-1],
          f"WACC: {report['wacc'] * 100:.2f}%   |   terminal g: "
          f"{report['dcf']['terminal_growth'] * 100:.1f}%\n\n"
          f"PV explicit FCFF:      Rs {bridge['pv_explicit']:,.0f} Cr\n"
          f"PV terminal value:     Rs {bridge['pv_terminal_gordon']:,.0f} Cr\n"
          f"Enterprise value:      Rs {bridge['enterprise_value']:,.0f} Cr\n"
          f"less net debt:         Rs {-bridge['net_debt']:,.0f} Cr\n"
          f"Equity value:          Rs {bridge['equity_value']:,.0f} Cr\n"
          f"Fair value per share:  Rs {per_share:,.2f}", size=15)

    # 6. Scenarios + sensitivity
    _slide(prs, "Scenarios and sensitivity",
           f"Bear/base/bull per share: "
           f"{report['scenarios']['bear']['per_share_inr']:,.0f} / "
           f"{report['scenarios']['base']['per_share_inr']:,.0f} / "
           f"{report['scenarios']['bull']['per_share_inr']:,.0f}.")
    if chart_paths and "sensitivity" in chart_paths:
        prs.slides[-1].shapes.add_picture(str(chart_paths["sensitivity"]),
                                          Inches(2.4), Inches(1.3), width=Inches(5.4))

    # 7. Key risks + DD questions
    _slide(prs, "Key risks and due-diligence questions",
           "Top flags per the rule engine; DD questions for management.")
    flag_lines = [f"[{flag['severity'].upper()}] {flag['rationale']}"
                  for flag in report["risk"]["flags"][:6]]
    flag_lines += ["", "DD questions:",
                   "- Customer A contract terms and renewal risk",
                   "- Expansion ramp schedule and margin recovery",
                   "- Refinancing plan for the FY2027 obligation"]
    _body(prs.slides[-1], "\n".join(flag_lines), top=1.2, size=12)

    prs.core_properties.title = f"FinSight AI — {company_name} — Investment Deck"
    prs.core_properties.author = "FinSight AI"
    prs.save(output_path)
    return output_path
