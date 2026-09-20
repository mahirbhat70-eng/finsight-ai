"""Excel model with LIVE FORMULAS (Playbook P6.3).

Sheets: Assumptions (yellow input cells per FAST convention, data
validation), Model (every computed cell is a formula referencing
Assumptions), Scenarios (3-column WACC x g), Checks (4 checks with
conditional formatting). NO hardcoded computed values in outputs — only
inputs and formulas.

Cell map is documented below and pinned by tests; the recalc test converts
via LibreOffice headless and asserts per-share == API result within 0.1%.
"""

import uuid
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.finmod.dcf import Forecast
from app.finmod.wacc import AssumptionSet

YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
MONEY = "#,##0.0"
PCT = "0.00"
SHARE = "#,##0.00"


def _wacc_inputs(series: dict, periods: list[str], assumptions: AssumptionSet
                 ) -> dict:
    t = len(periods) - 1

    def latest(key: str) -> float:
        return series[key][t]

    nwc = (latest("current_assets") - latest("cash")
           - latest("current_liabilities") + latest("st_borrowings"))
    return {
        "base_revenue": latest("revenue"),
        "nwc_days": nwc / latest("revenue") * 365.0,
        "base_nwc": nwc,
        "net_debt": latest("total_debt") - latest("cash"),
        "latest_ebitda": latest("ebitda"),
        "equity": latest("total_equity"),
        "total_debt": latest("total_debt"),
        "shares": latest("shares_outstanding"),
    }


def build_workbook(series: dict, periods: list[str],
                   assumptions: AssumptionSet, forecast: Forecast,
                   company_name: str = "NovaTech Industries Ltd",
                   assumptions_version: str = "api-1") -> Workbook:
    """The model is a formula network; inputs come from the statements."""
    wb = Workbook()
    inputs = _wacc_inputs(series, periods, assumptions)
    n = len(forecast.revenue_growth)

    # --- Assumptions -------------------------------------------------------
    ws = wb.active
    ws.title = "Assumptions"
    ws["A1"] = "INPUTS — yellow cells are editable (FAST convention)"
    ws["A1"].font = Font(bold=True, size=11)

    rows = [
        ("Risk-free rate (10Y G-Sec) %", assumptions.rf, PCT),
        ("Beta (levered)", assumptions.beta, "0.00"),
        ("Market risk premium %", assumptions.market_risk_premium, PCT),
        ("Pre-tax cost of debt %", assumptions.cost_of_debt_pre_tax, PCT),
        ("Tax rate %", assumptions.tax_rate, PCT),
        ("Mid-year convention (TRUE/FALSE)", assumptions.mid_year, None),
        ("Terminal growth g %", assumptions.terminal_growth, PCT),
        ("Exit EV/EBITDA multiple", assumptions.exit_multiple, "0.0"),
        ("Diluted shares (mn)", inputs["shares"], "#,##0"),
        ("Base revenue (INR cr)", inputs["base_revenue"], MONEY),
        ("NWC days (on revenue)", inputs["nwc_days"], "0.00"),
        ("Base NWC (INR cr)", inputs["base_nwc"], MONEY),
        ("Net debt (INR cr)", inputs["net_debt"], MONEY),
        ("Latest actual EBITDA (INR cr)", inputs["latest_ebitda"], MONEY),
        ("Equity, book proxy (INR cr)", inputs["equity"], MONEY),
        ("Total debt (INR cr)", inputs["total_debt"], MONEY),
    ]
    for r, (label, value, fmt) in enumerate(rows, start=2):
        ws.cell(row=r, column=1, value=label)
        cell = ws.cell(row=r, column=2, value=value)
        cell.fill = YELLOW
        if fmt:
            cell.number_format = fmt

    # computed WACC block (formulas, not values)
    ws["A19"] = "Cost of equity (CAPM) %"
    ws["B19"] = "=B2+B3*B4+0"  # + size premium if used; kept explicit
    ws["B19"].number_format = PCT
    ws["A20"] = "Cost of debt, after tax %"
    ws["B20"] = "=B5*(1-B6/100)"
    ws["B20"].number_format = PCT
    ws["A21"] = "WACC % (E/V x Ke + D/V x Kd_at)"
    ws["B21"] = "=(B16/(B16+B17))*B19+(B17/(B16+B17))*B20"
    ws["B21"].number_format = PCT
    ws["B21"].font = Font(bold=True)

    # forecast inputs
    ws["A23"] = "FORECAST — Y1..Y5 (yellow, editable)"
    ws["A23"].font = Font(bold=True, size=11)
    ws["A24"] = "Revenue growth %"
    ws["A25"] = "EBITDA margin %"
    ws["A26"] = "Capex % of revenue"
    ws["A27"] = "D&A % of revenue (single)"
    for i in range(n):
        col = 4 + i  # D..H
        header = ws.cell(row=23, column=col, value=f"Y{i + 1}")
        header.fill = HEADER_FILL
        header.font = HEADER_FONT
        for r, values in ((24, forecast.revenue_growth),
                          (25, forecast.ebitda_margin),
                          (26, forecast.capex_pct_revenue)):
            cell = ws.cell(row=r, column=col, value=values[i])
            cell.fill = YELLOW
            cell.number_format = "0.0"
    dna_cell = ws.cell(row=27, column=2, value=forecast.dna_pct_revenue)
    dna_cell.fill = YELLOW
    dna_cell.number_format = "0.0"
    ws["A29"] = f"{company_name} | assumptions version {assumptions_version}"
    ws["A30"] = "Generated by FinSight AI — deterministic engine; LLM never computes."

    ws.column_dimensions["A"].width = 36
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 13

    # data validation (playbook ranges: growth 0-50, WACC 1-60, tax 0-60)
    for formula, cells in (
        ('decimal 0 50', [f"D{i}:H{i}" for i in (24, 25, 26)]),
        ("decimal 0 60", ["B2", "B4", "B5", "B6", "B8"]),
        ("decimal 0.1 3", ["B3"]),
        ("decimal 1 60", ["B21"]),  # guardrail echo (B21 is a formula; validation is informational)
        ("decimal 1 30", ["B9"]),
    ):
        dv = DataValidation(type="decimal", formula1=formula.split()[1],
                            formula2=formula.split()[2], allow_blank=False,
                            showErrorMessage=True,
                            error="Value outside the allowed modelling range")
        ws.add_data_validation(dv)
        for rng in cells:
            dv.add(rng)  # string ranges: "B2" or "D24:H24"

    # --- Model -------------------------------------------------------------
    model = wb.create_sheet("Model")
    headers = ["Year", "Revenue", "EBITDA", "D&A", "EBIT", "NOPAT", "Capex",
               "NWC", "Delta NWC", "FCFF", "DF", "PV"]
    for c, h in enumerate(headers, start=1):
        cell = model.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for i in range(n):
        r = i + 2
        col = get_column_letter(4 + i)  # D..H on Assumptions
        model.cell(row=r, column=1, value=i + 1)
        if i == 0:
            model.cell(row=r, column=2, value="=Assumptions!$B$11*(1+Assumptions!D$24/100)")
            model.cell(row=r, column=9, value="=H2-Assumptions!$B$13")
        else:
            model.cell(row=r, column=2, value=f"=B{r - 1}*(1+Assumptions!{col}$24/100)")
            model.cell(row=r, column=9, value=f"=H{r}-H{r - 1}")
        model.cell(row=r, column=3, value=f"=B{r}*Assumptions!{col}$25/100")
        model.cell(row=r, column=4, value=f"=B{r}*Assumptions!$B$27/100")
        model.cell(row=r, column=5, value=f"=C{r}-D{r}")
        model.cell(row=r, column=6, value=f"=E{r}*(1-Assumptions!$B$6/100)")
        model.cell(row=r, column=7, value=f"=B{r}*Assumptions!{col}$26/100")
        model.cell(row=r, column=8, value=f"=B{r}*Assumptions!$B$12/365")
        model.cell(row=r, column=10, value=f"=F{r}+D{r}-G{r}-I{r}")
        model.cell(row=r, column=11,
                   value=f"=1/(1+Assumptions!$B$21/100)^(A{r}-IF(Assumptions!$B$7,0.5,0))")
        model.cell(row=r, column=12, value=f"=J{r}*K{r}")
        for c in (2, 3, 4, 5, 6, 7, 8, 9, 10, 12):
            model.cell(row=r, column=c).number_format = MONEY
        model.cell(row=r, column=11).number_format = "0.0000"

    dcf_rows = [
        ("PV of explicit FCFF", "=SUM(L2:L6)", MONEY),
        ("Terminal value (Gordon)", None, MONEY),
        ("PV of terminal value", None, MONEY),
        ("Enterprise value", None, MONEY),
        ("less: net debt", "=-Assumptions!$B$14", MONEY),
        ("Equity value", None, MONEY),
        ("Per share (INR)", None, SHARE),
    ]
    for r, (label, formula, fmt) in enumerate(dcf_rows, start=8):
        model.cell(row=r, column=1, value=label).font = Font(bold=(r in (11, 14)))
        if formula:
            model.cell(row=r, column=2, value=formula)
    model["B9"] = "=J6*(1+Assumptions!$B$8/100)/((Assumptions!$B$21-Assumptions!$B$8)/100)"
    model["B10"] = "=B9/(1+Assumptions!$B$21/100)^5"
    model["B11"] = "=B8+B10"
    model["B13"] = "=B11+B12"
    model["B14"] = "=B13*10/Assumptions!$B$10"
    model["A16"] = "Exit-multiple cross-check"
    model["A16"].font = Font(italic=True)
    model["B16"] = "=C6*Assumptions!$B$9"
    model["B17"] = "=B8+B16/(1+Assumptions!$B$21/100)^5"
    model["A17"] = "EV, exit method"
    model["A18"] = "Spread vs Gordon"
    model["B18"] = "=B17/B11-1"
    for addr in ("B8", "B9", "B10", "B11", "B12", "B13", "B16", "B17"):
        model[addr].number_format = MONEY
    model["B14"].number_format = SHARE
    model["B18"].number_format = "0.0%"
    model.column_dimensions["A"].width = 28
    for c in "BCDEFGHIJKL":
        model.column_dimensions[c].width = 13
    model.freeze_panes = "B2"

    # --- Scenarios ---------------------------------------------------------
    scen = wb.create_sheet("Scenarios")
    scen["A1"] = "Scenario"
    for c, name in enumerate(("Bear", "Base", "Bull"), start=2):
        cell = scen.cell(row=1, column=c, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    scen["A2"] = "WACC %"
    scen["B2"] = "=Assumptions!$B$21+1.5"
    scen["C2"] = "=Assumptions!$B$21"
    scen["D2"] = "=Assumptions!$B$21-1.5"
    scen["A3"] = "Terminal g %"
    scen["B3"] = "=Assumptions!$B$8-1"
    scen["C3"] = "=Assumptions!$B$8"
    scen["D3"] = "=Assumptions!$B$8+1"
    scen["A4"] = "Per share (INR)"
    for c in ("B", "C", "D"):
        scen[f"{c}4"] = (
            f"=(SUMPRODUCT(Model!$J$2:$J$6,1/(1+{c}$2/100)^(Model!$A$2:$A$6-0.5))"
            f"+Model!$J$6*(1+{c}$3/100)/(({c}$2-{c}$3)/100)/(1+{c}$2/100)^5"
            f"-Assumptions!$B$14)*10/Assumptions!$B$10")
        scen[f"{c}4"].number_format = SHARE
    scen["A6"] = ("Bear/bull growth and margin deltas live in the API engine; "
                  "flex them here by editing the yellow forecast cells.")
    scen.column_dimensions["A"].width = 30
    for c in "BCD":
        scen.column_dimensions[c].width = 16

    # --- Checks ------------------------------------------------------------
    checks = wb.create_sheet("Checks")
    checks["A1"] = "Model checks (conditional formatting: green PASS, red otherwise)"
    checks["A1"].font = Font(bold=True, size=11)
    check_rows = [
        ("EV identity: PV explicit + PV terminal = EV",
         "=IF(ABS(Model!B8+Model!B10-Model!B11)<0.01,\"PASS\",\"FAIL\")"),
        ("Terminal value share of EV below 75%",
         "=IF(Model!B10/Model!B11<=0.75,\"PASS\",\"WARN\")"),
        ("Discount factors strictly decreasing",
         "=IF(AND(Model!K2>Model!K3,Model!K3>Model!K4,Model!K4>Model!K5,"
         "Model!K5>Model!K6),\"PASS\",\"FAIL\")"),
        ("Implied EV/EBITDA within 3x-25x",
         "=IF(AND(Model!B11/Assumptions!$B$15>=3,Model!B11/Assumptions!$B$15<=25),"
         "\"PASS\",\"WARN\")"),
    ]
    for r, (label, formula) in enumerate(check_rows, start=3):
        checks.cell(row=r, column=1, value=label)
        checks.cell(row=r, column=2, value=formula)
        checks.cell(row=r, column=2).alignment = Alignment(horizontal="center")
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    checks.conditional_formatting.add(
        "B3:B6", CellIsRule(operator="equal", formula=['"PASS"'], fill=green))
    checks.conditional_formatting.add(
        "B3:B6", CellIsRule(operator="notEqual", formula=['"PASS"'], fill=red))
    checks.column_dimensions["A"].width = 48
    checks.column_dimensions["B"].width = 12

    return wb


def save_workbook(wb: Workbook, company_id: uuid.UUID, out_root: Path) -> Path:
    out_dir = out_root / str(company_id)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{uuid.uuid4().hex}.xlsx"
    wb.save(path)
    return path
