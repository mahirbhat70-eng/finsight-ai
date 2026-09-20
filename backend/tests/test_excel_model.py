"""Excel model tests (Playbook P6.3): live formulas, recalc via LibreOffice
headless, per-share equals the API/engine result within 0.1% for the base
case AND a flexed scenario (growth +2pp edited in the workbook).
"""

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def _find_fixtures() -> Path:
    for parent in [Path(__file__).resolve()] + list(Path(__file__).resolve().parents):
        cand = parent / "data" / "fixtures"
        if (cand / "novatech_statements.json").exists():
            return cand
    raise FileNotFoundError


FIXTURES = _find_fixtures()
SERIES = json.loads((FIXTURES / "novatech_statements.json").read_text())
SER, PERIODS = SERIES["items"], SERIES["periods"]

from app.exports.excel_model import build_workbook  # noqa: E402
from app.finmod.dcf import Forecast, run_dcf  # noqa: E402
from app.finmod.engine import run_company  # noqa: E402
from app.finmod.wacc import AssumptionSet  # noqa: E402


def test_every_computed_cell_is_a_formula() -> None:
    wb = build_workbook(SER, PERIODS, AssumptionSet(), Forecast())
    model = wb["Model"]
    for row in model.iter_rows(min_row=2, max_row=6, min_col=2, max_col=12):
        for cell in row:
            assert isinstance(cell.value, str) and cell.value.startswith("="), \
                f"Model!{cell.coordinate} must be a formula, got {cell.value!r}"
    for addr in ("B8", "B9", "B10", "B11", "B12", "B13", "B14"):
        assert str(model[addr].value).startswith("=")
    checks = wb["Checks"]
    for r in range(3, 7):
        assert str(checks.cell(row=r, column=2).value).startswith("=")

    assumptions = wb["Assumptions"]
    # inputs are values (not formulas); WACC block is formulas
    assert isinstance(assumptions["B2"].value, float)
    assert str(assumptions["B21"].value).startswith("=")
    # yellow input convention
    assert assumptions["B2"].fill.fgColor.rgb.endswith("FFF2CC")


def test_scenario_and_check_sheets_reference_model() -> None:
    wb = build_workbook(SER, PERIODS, AssumptionSet(), Forecast())
    scen = wb["Scenarios"]
    for cell in ("B4", "C4", "D4"):
        formula = str(scen[cell].value)
        assert "SUMPRODUCT" in formula and "Model!" in formula


@pytest.mark.skipif(shutil.which("soffice") is None,
                    reason="LibreOffice not installed")
@pytest.mark.parametrize("growth_delta", [0.0, 2.0])
def test_libreoffice_recalc_matches_engine(growth_delta: float) -> None:
    """Recalculate the workbook headless; per-share must match run_dcf
    within 0.1% for the base case and the flexed growth +2pp case."""
    assumptions = AssumptionSet()
    # build with the BASE forecast; the flex is applied to the yellow cells
    wb = build_workbook(SER, PERIODS, assumptions, Forecast())
    forecast = Forecast()
    if growth_delta:
        from openpyxl.styles import PatternFill

        ws = wb["Assumptions"]
        for i in range(5):
            col = chr(ord("D") + i)
            ws[f"{col}24"] = float(ws[f"{col}24"].value) + growth_delta
            ws[f"{col}24"].fill = PatternFill("solid", fgColor="FFF2CC")
        forecast = Forecast(
            revenue_growth=[g + growth_delta for g in forecast.revenue_growth])

    expected = run_dcf(SER, PERIODS, assumptions, forecast)["bridge"]["per_share_inr"]

    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "model.xlsx"
        wb.save(src)
        subprocess.run(  # noqa: S603
            ["soffice", "--headless", "--convert-to", "xlsx", "--outdir",
             str(Path(tmp) / "out"), str(src)],
            check=True, capture_output=True, timeout=120)
        converted = Path(tmp) / "out" / "model.xlsx"
        assert converted.exists(), "LibreOffice conversion produced no output"

        from openpyxl import load_workbook

        calc = load_workbook(converted, data_only=True)
        per_share = calc["Model"]["B14"].value
        wacc = calc["Assumptions"]["B21"].value
        checks = [calc["Checks"].cell(row=r, column=2).value for r in range(3, 7)]

    assert per_share is not None, "recalc produced no per-share value"
    assert abs(per_share - expected) / expected <= 0.001, (
        f"xlsx per-share {per_share} != engine {expected}")
    engine_wacc = run_company(SER, PERIODS)["wacc"] * 100
    assert abs(wacc - engine_wacc) <= 0.01
    # checks: identity PASS, TV-share may legitimately WARN on the base case
    # (82% of EV), DF monotonicity PASS, implied multiple PASS
    assert checks[0] == "PASS" and checks[2] == "PASS" and checks[3] == "PASS"
    assert checks[1] in ("PASS", "WARN"), f"unexpected TV check: {checks[1]}"
