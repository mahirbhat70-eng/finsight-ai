"""Analyst input-format xlsx fixture (Playbook P1.5): sheets Assumptions,
P&L, Balance Sheet, Cash Flow. Values only — the live-formula MODEL is the
Phase 6 export (app/exports/excel_model.py), not this fixture.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from novatech_model import build_model

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
SECTION_FONT = Font(bold=True, size=11)
NUM_FMT = "#,##0.0"


def build_xlsx(output_path: str, model: dict) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "Assumptions"
    rows = [
        ("Company", model["company"]),
        ("Currency", "INR crore unless stated"),
        ("Fiscal year end", "March"),
        ("Periods", ", ".join(model["periods"])),
        ("Tax rate (combined effective)", 0.2517),
        ("Average interest rate on debt", 0.092),
        ("Dividend payout policy", 0.20),
        ("Source", "Deterministic generator (seed 42); articulation asserted"),
    ]
    for r, (k, v) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=k).font = SECTION_FONT
        ws.cell(row=r, column=2, value=v)
        if isinstance(v, float):
            ws.cell(row=r, column=2).number_format = "0.00%"
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 70

    def sheet(name: str, rows: list[tuple[str, str]]) -> None:
        s = wb.create_sheet(name)
        s.cell(row=1, column=1, value="Rs in crore").font = SECTION_FONT
        for c, period in enumerate(model["periods"], start=2):
            cell = s.cell(row=1, column=c, value=period)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="right")
        for r, (label, key) in enumerate(rows, start=2):
            s.cell(row=r, column=1, value=label)
            for c, v in enumerate(model["items"][key], start=2):
                cell = s.cell(row=r, column=c, value=round(v, 1))
                cell.number_format = NUM_FMT
        s.column_dimensions["A"].width = 46
        for c in range(2, 2 + len(model["periods"])):
            s.column_dimensions[get_column_letter(c)].width = 13
        s.freeze_panes = "B2"

    from novatech_pdf import BS_ROWS, CF_ROWS, P_AND_L_ROWS

    sheet("P&L", [(lbl, key) for lbl, key in P_AND_L_ROWS if key in model["items"]])
    sheet("Balance Sheet", [(lbl, key) for lbl, key in BS_ROWS if key in model["items"]])
    sheet("Cash Flow", CF_ROWS)

    wb.save(output_path)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "novatech.xlsx"
    build_xlsx(out, build_model())
    print(f"wrote {out}")
